"""Lê o texto dos documentos da sala num processo separado (decisão D7 do plano do B-20).

O documento é de fora e pode ter sido feito para travar ou explorar o leitor. Por isso a leitura
roda num **subprocesso** do Python, com:
- tempo máximo (`TEMPO_MAXIMO_EM_SEGUNDOS`) e, no Linux, memória máxima (`RLIMIT_AS`), tempo de
  CPU e proibição de gravar arquivo (`RLIMIT_FSIZE = 0`), postos pelo próprio filho antes de
  tocar no documento; e a nota de "mate este primeiro" (`oom_score_adj`), para que, se a memória
  do contêiner acabar, quem cai é a leitura, não a API;
- o Python em modo isolado (`-I`) e um ambiente sem nenhuma variável do servidor: o filho não
  vê a URL do banco, a chave-mestra nem as chaves do Stripe;
- o documento já decifrado pela entrada padrão e o texto de volta em JSON pela saída padrão.

Falhou, estourou o tempo ou devolveu lixo: o documento fica `erro_na_leitura` e a API segue.
O filho confere de novo, pelo PDFium, o que a conferência de `arquivos.py` só vê por fora:
número de páginas, JavaScript e anexos.

Nenhum modelo de linguagem lê o documento (D9): o texto lido é dado, nunca instrução.
"""

import contextlib
import io
import json
import os
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

from src.sala import arquivos

VERSAO_DO_LEITOR = "sala-leitor-1"
TEMPO_MAXIMO_EM_SEGUNDOS = 60
MEMORIA_MAXIMA_EM_BYTES = 1024 * 1024 * 1024  # 1 GB
NOTA_DE_SACRIFICIO = (
    1000  # oom_score_adj máximo: o primeiro que o Linux mata sem memória
)

RAIZ_DO_PROJETO = Path(__file__).resolve().parents[2]
# variáveis que o Python precisa para abrir (no Windows, sem SYSTEMROOT ele nem inicia)
# SYSTEMROOT e SYSTEMDRIVE só existem no Windows (sem elas, o sistema grava um cache numa
# pasta chamada "%SystemDrive%" na pasta atual)
_VARIAVEIS_PASSADAS_AO_FILHO = ("PATH", "SYSTEMROOT", "SYSTEMDRIVE", "TEMP", "TMP")

_NS_DO_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass(frozen=True)
class Leitura:
    """O que a leitura achou: `situacao` é `lido`, `sem_texto` ou `erro_na_leitura`."""

    situacao: str
    paginas: tuple[str, ...] = ()
    motivo: str | None = (
        None  # código de `arquivos.MOTIVOS_DE_RECUSA`, quando o filho recusou
    )


ERRO = Leitura("erro_na_leitura")


# --- lado da API: dispara o filho e interpreta a resposta ----------------------------------


def _ambiente_do_filho() -> dict[str, str]:
    return {
        nome: os.environ[nome]
        for nome in _VARIAVEIS_PASSADAS_AO_FILHO
        if nome in os.environ
    }


def _comando_do_filho(tipo: str) -> list[str]:
    # -I: ignora PYTHONPATH e afins e não põe a pasta atual no caminho; a raiz do projeto entra
    # por argumento, e só ela
    codigo = (
        "import sys; sys.path.insert(0, sys.argv[1]); "
        "from src.sala.leitura import principal; principal(sys.argv[2])"
    )
    return [sys.executable, "-I", "-c", codigo, str(RAIZ_DO_PROJETO), tipo]


def _interpretar(saida: bytes) -> Leitura:
    try:
        resposta = json.loads(saida)
    except (ValueError, UnicodeDecodeError):
        return ERRO
    if not isinstance(resposta, dict):
        return ERRO
    motivo = resposta.get("recusa")
    if motivo is not None:
        return Leitura(
            "erro_na_leitura",
            motivo=motivo if motivo in arquivos.MOTIVOS_DE_RECUSA else None,
        )
    paginas = resposta.get("paginas")
    if not isinstance(paginas, list) or not all(isinstance(p, str) for p in paginas):
        return ERRO
    tem_texto = any(pagina.strip() for pagina in paginas)
    return Leitura("lido" if tem_texto else "sem_texto", tuple(paginas))


def ler_em_subprocesso(
    conteudo: bytes, tipo: str, tempo_maximo: float = TEMPO_MAXIMO_EM_SEGUNDOS
) -> Leitura:
    """O texto do documento por página, lido num processo separado e com limites."""
    if tipo not in arquivos.TIPOS_ACEITOS:
        raise ValueError(f"Tipo de documento: {arquivos.TIPOS_ACEITOS}.")
    # o filho trabalha numa pasta vazia que some depois: nada do que ele (ou o sistema, por ele)
    # gravar cai na pasta do projeto
    with tempfile.TemporaryDirectory(prefix="leitura-da-sala-") as pasta_do_filho:
        try:
            resultado = subprocess.run(
                _comando_do_filho(tipo),
                input=conteudo,
                capture_output=True,
                timeout=tempo_maximo,
                env=_ambiente_do_filho(),
                cwd=pasta_do_filho,
                check=False,
            )
        except (subprocess.TimeoutExpired, OSError):
            return ERRO  # subprocess.run já matou o filho
    if resultado.returncode != 0:
        return ERRO
    return _interpretar(resultado.stdout)


# --- lado do filho: põe os limites, lê e responde -----------------------------------------


def _limitar_o_proprio_processo() -> None:
    """No Linux, o filho limita a si mesmo antes de abrir o documento (memória, CPU, arquivos)."""
    if not sys.platform.startswith("linux"):
        return
    import resource  # só existe em sistemas Unix

    resource.setrlimit(
        resource.RLIMIT_AS, (MEMORIA_MAXIMA_EM_BYTES, MEMORIA_MAXIMA_EM_BYTES)
    )
    resource.setrlimit(
        resource.RLIMIT_CPU, (TEMPO_MAXIMO_EM_SEGUNDOS, TEMPO_MAXIMO_EM_SEGUNDOS)
    )
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))  # não grava arquivo nenhum
    with contextlib.suppress(
        OSError
    ):  # sem /proc (alguns contêineres): ficam os outros limites
        Path("/proc/self/oom_score_adj").write_text(str(NOTA_DE_SACRIFICIO))


def _paginas_do_pdf(conteudo: bytes) -> list[str]:
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    try:
        documento = pdfium.PdfDocument(conteudo)
    except pdfium.PdfiumError as erro:
        com_senha = getattr(erro, "err_code", None) == pdfium_c.FPDF_ERR_PASSWORD
        motivo = arquivos.PDF_CIFRADO if com_senha else arquivos.ARQUIVO_DANIFICADO
        raise arquivos.ArquivoRecusado(motivo) from None
    try:
        if pdfium_c.FPDF_GetSecurityHandlerRevision(documento.raw) != -1:
            raise arquivos.ArquivoRecusado(arquivos.PDF_CIFRADO)
        if (
            pdfium_c.FPDFDoc_GetJavaScriptActionCount(documento.raw) > 0
            or pdfium_c.FPDFDoc_GetAttachmentCount(documento.raw) > 0
        ):
            raise arquivos.ArquivoRecusado(arquivos.PDF_ATIVO)
        if len(documento) > arquivos.PAGINAS_MAXIMAS:
            raise arquivos.ArquivoRecusado(arquivos.PAGINAS_DEMAIS)
        return [pagina.get_textpage().get_text_bounded() for pagina in documento]
    finally:
        documento.close()


def _paginas_do_xlsx(conteudo: bytes) -> list[str]:
    import openpyxl

    # read_only e data_only: lê os valores gravados, sem avaliar fórmula nem carregar macro
    livro = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        if len(livro.worksheets) > arquivos.PAGINAS_MAXIMAS:
            raise arquivos.ArquivoRecusado(arquivos.PAGINAS_DEMAIS)
        return [
            "\n".join(
                "\t".join("" if valor is None else str(valor) for valor in linha)
                for linha in aba.iter_rows(values_only=True)
            )
            for aba in livro.worksheets
        ]
    finally:
        livro.close()


def _paginas_do_docx(conteudo: bytes) -> list[str]:
    import zipfile

    from defusedxml import ElementTree

    with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
        raiz = ElementTree.fromstring(pacote.read("word/document.xml"))
    paginas, linhas, linha = [], [], []
    for elemento in raiz.iter():
        if elemento.tag == f"{_NS_DO_WORD}t":
            linha.append(elemento.text or "")
        elif elemento.tag == f"{_NS_DO_WORD}tab":
            linha.append("\t")
        elif elemento.tag == f"{_NS_DO_WORD}p":
            linhas.append("".join(linha))
            linha = []
        elif _quebra_de_pagina(elemento.tag, elemento.get(f"{_NS_DO_WORD}type")):
            linhas.append("".join(linha))
            paginas.append("\n".join(linhas).strip("\n"))
            linhas, linha = [], []
    linhas.append("".join(linha))
    paginas.append("\n".join(linhas).strip("\n"))
    return paginas


def _quebra_de_pagina(etiqueta: str, tipo: str | None) -> bool:
    # o DOCX não tem página: vale a quebra manual e a última quebra que o Word desenhou
    return etiqueta == f"{_NS_DO_WORD}lastRenderedPageBreak" or (
        etiqueta == f"{_NS_DO_WORD}br" and tipo == "page"
    )


def _paginas_do_csv(conteudo: bytes) -> list[str]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252")
    return [texto]


_LEITORES = {
    "pdf": _paginas_do_pdf,
    "xlsx": _paginas_do_xlsx,
    "docx": _paginas_do_docx,
    "csv": _paginas_do_csv,
}


def principal(tipo: str) -> None:
    """Entrada do filho: lê o documento da entrada padrão e escreve o JSON na saída padrão."""
    _limitar_o_proprio_processo()
    conteudo = sys.stdin.buffer.read()
    try:
        resposta: dict[str, object] = {"paginas": _LEITORES[tipo](conteudo)}
    except arquivos.ArquivoRecusado as recusa:
        resposta = {"recusa": recusa.motivo}
    sys.stdout.write(json.dumps(resposta, ensure_ascii=True))
    sys.stdout.flush()
