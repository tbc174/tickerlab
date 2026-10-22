"""O deck (PDF e PPT) da ficha e do relatório do setor: os mesmos slides nos dois formatos, os números
da tela, a fonte em cada número e "Limites e cuidados" sempre no fim."""

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from pptx import Presentation
from pypdf import PdfReader

from src import apresentacao as ap
from src import dossie
from src import relatorio_empresa as re_
from src import relatorio_setor as rs
from src.apresentacao_pdf import apresentacao_em_pdf
from src.apresentacao_pptx import apresentacao_em_pptx
from src.estilos_pdf import ESTILOS
from tests.test_relatorio_empresa import _entradas as entradas_da_empresa
from tests.test_relatorio_empresa import dados_da_weg  # noqa: F401
from tests.test_relatorio_setor import HOJE, gravar_noite
from tests.test_relatorio_setor import _entradas as entradas_do_setor

GERADO_EM = datetime(2026, 10, 22, 9, 0, tzinfo=ZoneInfo("America/Sao_Paulo"))


def _sem_espacos(texto: str) -> str:
    return "".join(texto.split())


def _texto_do_pdf(pdf: bytes) -> str:
    return " ".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)


def _texto_do_pptx(pptx: bytes) -> str:
    textos = []
    for slide in Presentation(BytesIO(pptx)).slides:
        for forma in slide.shapes:
            if forma.has_text_frame:
                textos.append(forma.text_frame.text)
            if forma.has_table:
                textos += [c.text for linha in forma.table.rows for c in linha.cells]
    return " ".join(textos)


@pytest.fixture
def ficha(dados_da_weg):  # noqa: F811
    return re_.montar(entradas_da_empresa(dados_da_weg))


@pytest.fixture
def setor(tmp_path):
    secao = dossie.secao_do_setor("eletricas", gravar_noite(tmp_path))
    assert secao is not None
    return rs.montar_relatorio_do_setor(entradas_do_setor(secao), "pro", HOJE)


def test_ficha_mesmos_slides_no_pdf_e_no_ppt(ficha) -> None:
    deck = ap.da_ficha(ficha, GERADO_EM)
    pdf = apresentacao_em_pdf(deck, GERADO_EM)
    pptx = apresentacao_em_pptx(deck, GERADO_EM)
    assert pdf.startswith(b"%PDF") and pptx[:2] == b"PK"
    assert len(PdfReader(BytesIO(pdf)).pages) == len(deck.slides)
    assert len(Presentation(BytesIO(pptx)).slides) == len(deck.slides)


def test_ficha_tem_os_numeros_da_tela_com_a_fonte(ficha) -> None:
    deck = ap.da_ficha(ficha, GERADO_EM)
    numeros = [s for s in deck.slides if s.tipo == ap.NUMEROS]
    assert numeros
    receita = next(i for i in ficha.secao("resultado").itens if i.chave == "receita")
    linha = next(
        linha for s in numeros for linha in s.linhas if linha[0] == receita.nome
    )
    assert linha[1] == receita.texto and linha[3] == (receita.fonte or "—")
    for texto in (
        _texto_do_pdf(apresentacao_em_pdf(deck, GERADO_EM)),
        _texto_do_pptx(apresentacao_em_pptx(deck, GERADO_EM)),
    ):
        assert _sem_espacos(receita.texto) in _sem_espacos(texto)


def test_ficha_fecha_com_fontes_e_limites(ficha) -> None:
    deck = ap.da_ficha(ficha, GERADO_EM)
    assert deck.slides[-1].tipo == ap.LIMITES
    assert deck.slides[-1].titulo.startswith(ap.TITULO_DOS_LIMITES)
    assert any(s.tipo == ap.FONTES for s in deck.slides)
    assert "recomendação de investimento" in " ".join(deck.textos())
    limites = _sem_espacos(ficha.secao("limites").paragrafos[0])
    assert limites in _sem_espacos(_texto_do_pdf(apresentacao_em_pdf(deck, GERADO_EM)))


@pytest.mark.parametrize("plano", ["pro", "gratuito", None])
def test_setor_tem_limites_em_qualquer_plano_e_os_mesmos_slides(
    da_noite_plano, plano
) -> None:
    relatorio = rs.montar_relatorio_do_setor(da_noite_plano, plano, HOJE)
    deck = ap.do_setor(relatorio, GERADO_EM)
    assert deck.slides[-1].tipo == ap.LIMITES
    pdf = apresentacao_em_pdf(deck, GERADO_EM)
    pptx = apresentacao_em_pptx(deck, GERADO_EM)
    assert len(PdfReader(BytesIO(pdf)).pages) == len(deck.slides)
    assert len(Presentation(BytesIO(pptx)).slides) == len(deck.slides)
    assert "mercado relevante do CADE" in _texto_do_pdf(pdf)
    assert "mercado relevante do CADE" in _texto_do_pptx(pptx)


@pytest.fixture
def da_noite_plano(tmp_path):
    secao = dossie.secao_do_setor("eletricas", gravar_noite(tmp_path))
    assert secao is not None
    return entradas_do_setor(secao)


def test_setor_tem_as_secoes_da_tela_na_ordem_e_as_fontes(setor) -> None:
    deck = ap.do_setor(setor, GERADO_EM)
    titulos = " | ".join(s.titulo for s in deck.slides)
    posicao = 0
    for secao in setor.secoes:
        posicao = titulos.index(secao.titulo, posicao) + 1
    fontes = next(s for s in deck.slides if s.tipo == ap.FONTES)
    assert set(setor.fontes) <= set(fontes.topicos)
    assert "CPFE3" in " ".join(deck.textos())  # sigla preservada nas tabelas


def test_tabela_larga_fica_com_a_primeira_coluna_e_as_mais_recentes() -> None:
    linhas = [{f"c{i}": f"v{i}" for i in range(10)}]
    cabecalho, corpo = ap._tabela_de_texto(linhas, str, str)
    assert len(cabecalho) == ap.COLUNAS_NA_TABELA
    assert cabecalho[0] == "c0" and cabecalho[-1] == "c9"
    assert corpo == [("v0", "v5", "v6", "v7", "v8", "v9")]


def test_texto_longo_vira_mais_de_um_slide() -> None:
    topicos = ["x" * 400, "y" * 400, "z" * 400]
    slides = ap._slides_de_texto("Titulo", topicos)
    assert len(slides) == 3 and slides[0].titulo == "Titulo (1/3)"


@pytest.mark.parametrize("chave", list(ESTILOS))
def test_os_tres_estilos_geram_o_pdf(ficha, chave) -> None:
    deck = ap.da_ficha(ficha, GERADO_EM)
    assert apresentacao_em_pdf(deck, GERADO_EM, ESTILOS[chave]).startswith(b"%PDF")
