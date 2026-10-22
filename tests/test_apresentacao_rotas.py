"""As rotas do deck (PDF e PPT) da ficha e do setor: mesmo acesso e mesmo plano das rotas da tela,
o arquivo certo para cada formato e "Limites e cuidados" em qualquer plano."""

from io import BytesIO

import pytest
from pptx import Presentation
from pypdf import PdfReader

from src.banco.usuarios import mudar_plano
from tests import test_api_v1_rotas as rotas
from tests.test_api_v1_rotas import _com_a_petrobras_coletada
from tests.test_relatorio_do_setor_pdf import com_a_noite  # noqa: F401

chave = rotas.chave
cliente = rotas.cliente
SETOR = "/api/v1/relatorio-setor/apresentacao"
FICHA = "/api/v1/relatorio-empresa/apresentacao"
TIPO_DO_PPT = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def _como_pro(banco) -> None:
    with banco.begin() as conexao:
        mudar_plano(conexao, rotas._id_da_ana_no(banco), "pro")


def _texto_do_ppt(conteudo: bytes) -> str:
    textos = []
    for slide in Presentation(BytesIO(conteudo)).slides:
        for forma in slide.shapes:
            if forma.has_text_frame:
                textos.append(forma.text_frame.text)
            if forma.has_table:
                textos += [c.text for linha in forma.table.rows for c in linha.cells]
    return " ".join(textos)


def test_setor_em_pdf_e_em_ppt_para_o_pro(cliente, banco, com_a_noite) -> None:  # noqa: F811
    _como_pro(banco)
    pdf = cliente.get(SETOR, params={"setor": "eletricas", "formato": "pdf"})
    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.headers["cache-control"] == "no-store"
    assert "setor-eletricas-tickerlab.pdf" in pdf.headers["content-disposition"]
    ppt = cliente.get(SETOR, params={"setor": "eletricas", "formato": "pptx"})
    assert ppt.status_code == 200, ppt.text
    assert ppt.headers["content-type"] == TIPO_DO_PPT
    assert "setor-eletricas-tickerlab.pptx" in ppt.headers["content-disposition"]
    paginas = len(PdfReader(BytesIO(pdf.content)).pages)
    assert paginas == len(Presentation(BytesIO(ppt.content)).slides)
    assert "Limites e cuidados" in _texto_do_ppt(ppt.content)


def test_setor_do_gratuito_vem_trancado_e_com_limites(cliente, com_a_noite) -> None:  # noqa: F811
    ppt = cliente.get(SETOR, params={"setor": "eletricas", "formato": "pptx"})
    assert ppt.status_code == 200
    texto = _texto_do_ppt(ppt.content)
    assert "HHI (pontos)" not in texto and "Limites e cuidados" in texto


@pytest.mark.parametrize("setor", ["aviacao", "../eletricas"])
def test_setor_desconhecido_404(cliente, com_a_noite, setor) -> None:  # noqa: F811
    resposta = cliente.get(SETOR, params={"setor": setor, "formato": "pdf"})
    assert resposta.status_code == 404
    assert setor not in resposta.text


def test_formato_fora_da_lista_e_sem_login(cliente) -> None:
    assert (
        cliente.get(SETOR, params={"setor": "eletricas", "formato": "doc"}).status_code
        == 422
    )
    assert cliente.get(SETOR, params={"setor": "eletricas"}).status_code == 422
    cliente.headers.pop("Cookie")
    resposta = cliente.get(SETOR, params={"setor": "eletricas", "formato": "pdf"})
    assert resposta.status_code == 401


def test_ficha_em_ppt_do_pro_tem_os_numeros_da_base(
    cliente, banco, tmp_path, monkeypatch
) -> None:
    _com_a_petrobras_coletada(banco, tmp_path, monkeypatch)
    _como_pro(banco)
    ppt = cliente.get(FICHA, params={"ticker": "petr4", "formato": "pptx"})
    assert ppt.status_code == 200, ppt.text
    assert "ficha-PETR4-tickerlab.pptx" in ppt.headers["content-disposition"]
    texto = _texto_do_ppt(ppt.content)
    assert "R$ 495,2 mi" in texto and "Limites e cuidados" in texto
    pdf = cliente.get(FICHA, params={"ticker": "petr4", "formato": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_ficha_do_gratuito_vem_trancada_sem_numeros(
    cliente, banco, tmp_path, monkeypatch
) -> None:
    _com_a_petrobras_coletada(banco, tmp_path, monkeypatch)
    ppt = cliente.get(FICHA, params={"ticker": "PETR4.SA", "formato": "pptx"})
    assert ppt.status_code == 200, ppt.text
    texto = _texto_do_ppt(ppt.content)
    assert "R$ 495,2 mi" not in texto and "PRICEWATERHOUSECOOPERS" not in texto
    assert "Limites e cuidados" in texto


def test_ficha_de_ticker_sem_empresa_da_404(cliente) -> None:
    resposta = cliente.get(FICHA, params={"ticker": "ZZZZ9.SA", "formato": "pdf"})
    assert resposta.status_code == 404
