"""Preços mensais de referência (Pink Sheet do Banco Mundial e FMI) no acervo de séries, sem
internet: os dublês de `dados_de_exemplo/commodities/` são 12 meses reais de cada fonte."""

import logging
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from src import precos_de_referencia as precos
from src.banco import series as acervo_de_series
from src.banco.coletas import Marca, licenca_da_fonte, marca_da_fonte

PASTA = Path(__file__).parent.parent / "dados_de_exemplo" / "commodities"
PINK = (PASTA / "pink_sheet_mensal.xlsx").read_bytes()
FMI = (PASTA / "fmi_pcps_mensal.csv").read_bytes()
SETEMBRO = date(2026, 9, 1)


def _por_id(valores: list[precos.ValorMensal]) -> dict[tuple[str, date], float | None]:
    return {(v.id, v.mes): v.valor for v in valores}


def _planilha_alterada(tmp_path: Path, alterar) -> bytes:
    """O dublê do Pink Sheet depois de `alterar(aba)`."""
    caminho = tmp_path / "alterada.xlsx"
    pasta = openpyxl.load_workbook(PASTA / "pink_sheet_mensal.xlsx")
    alterar(pasta.active)
    pasta.save(caminho)
    return caminho.read_bytes()


def _coluna(aba, nome: str) -> int:
    return next(c.column for c in aba[5] if c.value and nome in str(c.value))


# --- Pink Sheet ---------------------------------------------------------------------------------
def test_pink_sheet_traz_os_itens_com_mes_no_primeiro_dia():
    valores = _por_id(precos.ler_pink_sheet(PINK))
    assert valores[("bm-pink-cafe-arabica", SETEMBRO)] == pytest.approx(7.26)
    assert valores[("bm-pink-cafe-robusta", SETEMBRO)] == pytest.approx(3.69)
    assert valores[("bm-pink-ureia", SETEMBRO)] == pytest.approx(407.5)
    ids = {i for i, _ in valores}
    assert ids == {f"bm-pink-{item.chave}" for item in precos.ITENS}
    assert {m for _, m in valores} == {
        date(2025, 10, 1) if m == 10 else date(2025 + m // 13, (m - 1) % 12 + 1, 1)
        for m in range(10, 22)
    }


def test_unidades_dos_itens_ficam_no_catalogo():
    unidades = {item.chave: item.unidade for item in precos.ITENS}
    assert unidades["cafe-arabica"] == "US$/kg"
    assert unidades["ureia"] == "US$/t"
    assert precos.serie_do_acervo("bm-pink-cafe-arabica").unidade == "US$/kg"


def test_celula_vazia_vira_none_e_nao_zero(tmp_path):
    def esvaziar(aba):
        aba.cell(row=aba.max_row, column=_coluna(aba, "Coffee, Arabica")).value = ".."

    valores = _por_id(precos.ler_pink_sheet(_planilha_alterada(tmp_path, esvaziar)))
    assert valores[("bm-pink-cafe-arabica", SETEMBRO)] is None
    assert valores[("bm-pink-cafe-robusta", SETEMBRO)] == pytest.approx(3.69)


def test_item_que_sumiu_da_planilha_e_ignorado_com_aviso(tmp_path, caplog):
    def tirar_ureia(aba):
        aba.cell(row=5, column=_coluna(aba, "Urea")).value = "Urea, nova versão"

    with caplog.at_level(logging.WARNING):
        valores = precos.ler_pink_sheet(_planilha_alterada(tmp_path, tirar_ureia))
    assert "bm-pink-ureia" not in {v.id for v in valores}
    assert "Urea" in caplog.text
    assert "bm-pink-cafe-arabica" in {v.id for v in valores}


def test_unidade_diferente_na_planilha_ignora_o_item_com_aviso(tmp_path, caplog):
    def trocar_unidade(aba):
        aba.cell(row=6, column=_coluna(aba, "Coffee, Arabica")).value = "(c/lb)"

    with caplog.at_level(logging.WARNING):
        valores = precos.ler_pink_sheet(_planilha_alterada(tmp_path, trocar_unidade))
    assert "bm-pink-cafe-arabica" not in {v.id for v in valores}
    assert "unidade" in caplog.text


# --- FMI ----------------------------------------------------------------------------------------
def test_fmi_converte_para_a_unidade_do_pink_sheet():
    valores = _por_id(precos.ler_fmi(FMI))
    # 330,12 centavos de dólar por libra-peso → US$ por kg
    assert valores[("fmi-cafe-arabica", SETEMBRO)] == pytest.approx(7.278, rel=1e-3)
    assert valores[("fmi-ureia", SETEMBRO)] == pytest.approx(260.27, rel=1e-3)
    assert {i for i, _ in valores} == {f"fmi-{item.chave}" for item in precos.ITENS}


def test_fmi_ignora_indicador_desconhecido_com_aviso(caplog):
    extra = FMI + b"G001,PXYZ,USD,M,2026-M09,1.5\n"
    with caplog.at_level(logging.WARNING):
        valores = precos.ler_fmi(extra)
    assert "PXYZ" in caplog.text
    assert all("pxyz" not in v.id for v in valores)


def test_fmi_so_le_mensal_em_dolar_e_vazio_vira_none():
    corpo = (
        b"COUNTRY,INDICATOR,DATA_TRANSFORMATION,FREQUENCY,TIME_PERIOD,OBS_VALUE\n"
        b"G001,PUREA,USD,M,2026-M09,\n"
        b"G001,PUREA,INDEX,M,2026-M09,180.5\n"
        b"G001,PUREA,USD,Q,2026-Q3,250\n"
        b",,,,,\n"
    )
    assert _por_id(precos.ler_fmi(corpo)) == {("fmi-ureia", SETEMBRO): None}


def test_pink_sheet_e_fmi_concordam_em_1_por_cento_onde_a_fonte_e_a_mesma():
    pink = _por_id(precos.ler_pink_sheet(PINK))
    fmi = _por_id(precos.ler_fmi(FMI))
    assert precos.COMPARAVEIS == ("cafe-arabica", "cafe-robusta", "algodao")
    for chave in precos.COMPARAVEIS:
        for (codigo, mes), valor in pink.items():
            if codigo == f"bm-pink-{chave}":
                assert valor is not None
                assert fmi[(f"fmi-{chave}", mes)] == pytest.approx(valor, rel=0.01)


def test_fertilizantes_divergem_porque_os_indicadores_sao_outros():
    # honestidade com os números: ureia, cloreto de potássio e DAP não são o mesmo indicador nas
    # duas fontes; o FMI serve de conferência só para o que concorda (COMPARAVEIS)
    pink = _por_id(precos.ler_pink_sheet(PINK))
    fmi = _por_id(precos.ler_fmi(FMI))
    assert (
        abs(fmi[("fmi-ureia", SETEMBRO)] / pink[("bm-pink-ureia", SETEMBRO)] - 1) > 0.05
    )
    assert "ureia" not in precos.COMPARAVEIS


# --- acervo (como era na data) ------------------------------------------------------------------
def test_gravar_cadastra_a_serie_e_guarda_a_safra(banco):
    valores = precos.ler_pink_sheet(PINK)
    with banco.begin() as conexao:
        novas = precos.gravar_precos(conexao, valores, date(2026, 10, 3))
    assert novas == len(valores)
    with banco.connect() as conexao:
        serie = acervo_de_series.como_estava_em(
            conexao, "bm-pink-cafe-arabica", date(2026, 10, 3)
        )
    assert serie.iloc[-1] == pytest.approx(7.26)
    assert len(serie) == 12


def test_celula_vazia_nao_grava_zero(banco, tmp_path):
    def esvaziar(aba):
        aba.cell(row=aba.max_row, column=_coluna(aba, "Coffee, Arabica")).value = ".."

    valores = precos.ler_pink_sheet(_planilha_alterada(tmp_path, esvaziar))
    with banco.begin() as conexao:
        precos.gravar_precos(conexao, valores, date(2026, 10, 3))
    with banco.connect() as conexao:
        serie = acervo_de_series.como_estava_em(
            conexao, "bm-pink-cafe-arabica", date(2026, 10, 3)
        )
    assert len(serie) == 11
    assert pd_nao_tem_data(serie, SETEMBRO)


def pd_nao_tem_data(serie, dia: date) -> bool:
    return not any(d.date() == dia for d in serie.index)


def test_sem_olhar_o_futuro_revisao_nao_muda_o_passado(banco, tmp_path):
    def revisar(aba):
        aba.cell(row=aba.max_row, column=_coluna(aba, "Coffee, Arabica")).value = 7.5

    with banco.begin() as conexao:
        precos.gravar_precos(conexao, precos.ler_pink_sheet(PINK), date(2026, 10, 3))
    revisada = precos.ler_pink_sheet(_planilha_alterada(tmp_path, revisar))
    with banco.begin() as conexao:
        assert precos.gravar_precos(conexao, revisada, date(2026, 11, 3)) == 1
    with banco.connect() as conexao:
        antes = acervo_de_series.como_estava_em(
            conexao, "bm-pink-cafe-arabica", date(2026, 10, 31)
        )
        depois = acervo_de_series.como_estava_em(
            conexao, "bm-pink-cafe-arabica", date(2026, 11, 30)
        )
        todas = acervo_de_series.safras(conexao, "bm-pink-cafe-arabica", SETEMBRO)
    assert antes.iloc[-1] == pytest.approx(7.26)
    assert depois.iloc[-1] == pytest.approx(7.5)
    assert [safra for safra, _ in todas] == [date(2026, 10, 3), date(2026, 11, 3)]


def test_repetir_a_coleta_nao_duplica(banco):
    valores = precos.ler_pink_sheet(PINK)
    with banco.begin() as conexao:
        precos.gravar_precos(conexao, valores, date(2026, 10, 3))
    with banco.begin() as conexao:
        assert precos.gravar_precos(conexao, valores, date(2026, 10, 4)) == 0


# --- licença ------------------------------------------------------------------------------------
def test_series_de_referencia_ficam_so_para_uso_interno():
    for fonte in ("bm-pink-cafe-arabica", "fmi-cafe-arabica"):
        assert precos.serie_do_acervo(fonte).uso_interno
        assert marca_da_fonte(fonte) in {Marca.USO_INTERNO, Marca.A_CONFIRMAR}
    assert "CC BY 4.0" in licenca_da_fonte("bm-pink-ureia")
    assert "permissão" in licenca_da_fonte("fmi-ureia")


# --- endereços ----------------------------------------------------------------------------------
def test_acha_o_endereco_do_pink_sheet_na_pagina():
    html = (
        '<a href="https://thedocs.worldbank.org/en/doc/abc123-0050012026/related/'
        'CMO-Historical-Data-Monthly.xlsx">October 2026 (XLS)</a>'
    )
    assert precos.achar_endereco_do_pink(html).endswith(
        "abc123-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
    )
    with pytest.raises(precos.PrecosIndisponiveisError):
        precos.achar_endereco_do_pink("<html>sem link</html>")


def test_calendario_da_coleta():
    # o Pink Sheet sai nos primeiros dias úteis; o FMI, perto do dia 10; vazio, coleta já
    assert precos.deve_coletar("bm-pink", date(2026, 10, 3), tem_dados=True)
    assert not precos.deve_coletar("bm-pink", date(2026, 10, 20), tem_dados=True)
    assert precos.deve_coletar("bm-pink", date(2026, 10, 20), tem_dados=False)
    assert precos.deve_coletar("fmi-pcps", date(2026, 10, 10), tem_dados=True)
    assert not precos.deve_coletar("fmi-pcps", date(2026, 10, 3), tem_dados=True)
